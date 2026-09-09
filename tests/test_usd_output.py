from __future__ import annotations

import json
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from src.models import Transaction
from src.output_writer import (
    EXCEPTION_COLUMNS,
    RPA_BUSINESS_COLUMNS,
    RPA_CHI_TIEN_MAT_COLUMNS,
    RPA_THU_TIEN_MAT_COLUMNS,
    write_outputs,
)
from src.rpa_summary import MESSAGE_COLUMN, STATUS_COLUMN, SUMMARY_SHEET_NAME, VOUCHER_COLUMN, mark_rpa_done
from src.rpa_tracking import STATUS_DONE
from src.usd.adapter import outcomes_to_processed
from src.usd.processor import USDProcessor
from src.usd.profile import load_usd_profile


def _transaction(description, *, debit=0, credit=0, day=1, row=14):
    return Transaction(
        source_file="usd.xlsx",
        bank="MSB",
        transaction_date=date(2026, 6, day),
        doc_no=f"REF-{row}",
        description=description,
        counterparty_raw="",
        debit_amount=debit,
        credit_amount=credit,
        original_row_index=row,
        source_sheet="Sheet1",
    )


def _processed_usd_batch():
    body = json.dumps(
        {
            "currencyOverview": [
                {
                    "currencyCode": "USD",
                    "exchangeRatesData": [{"currencyMarket": 1, "buyRateValue": "26,120.125"}],
                }
            ]
        }
    )

    def transport(url, timeout):
        return (503, "") if "02%2F06%2F2026" in url else (200, body)

    transactions = [
        _transaction("RTM NGUYEN VAN AN CCCD 1", debit=0.125, row=14),
        _transaction("THU TIEN TAU", credit=2.5, row=15),
        _transaction("PHI CHUYEN TIEN", debit=0.5, row=16),
        _transaction("THANH TOAN TAU", debit=3.75, day=2, row=17),
    ]
    outcomes = USDProcessor(load_usd_profile("config/usd_msb.yaml"), transport=transport).process_batch(transactions)
    return outcomes_to_processed(outcomes)


def _headers(ws):
    return [cell.value for cell in ws[1]]


def test_usd_writer_keeps_five_sheet_schema_and_routes_skip_and_exception(tmp_path):
    processed = _processed_usd_batch()
    result = write_outputs(processed, tmp_path, {"output": {}})
    workbook = load_workbook(result.excel_path, data_only=True)

    assert workbook.sheetnames == [
        "BAO_NO_INPUT",
        "BAO_CO_INPUT",
        "THU_TIEN_MAT_INPUT",
        "CHI_TIEN_MAT_INPUT",
        "EXCEPTION",
    ]
    assert _headers(workbook["BAO_NO_INPUT"]) == RPA_BUSINESS_COLUMNS
    assert _headers(workbook["BAO_CO_INPUT"]) == RPA_BUSINESS_COLUMNS
    assert _headers(workbook["THU_TIEN_MAT_INPUT"]) == RPA_THU_TIEN_MAT_COLUMNS
    assert _headers(workbook["CHI_TIEN_MAT_INPUT"]) == RPA_CHI_TIEN_MAT_COLUMNS
    assert _headers(workbook["EXCEPTION"]) == EXCEPTION_COLUMNS

    assert workbook["BAO_NO_INPUT"].max_row == 1
    assert workbook["BAO_CO_INPUT"].max_row == 2
    assert workbook["THU_TIEN_MAT_INPUT"].max_row == 2
    assert workbook["CHI_TIEN_MAT_INPUT"].max_row == 2
    assert workbook["EXCEPTION"].max_row == 2

    receipt = workbook["THU_TIEN_MAT_INPUT"]
    payment = workbook["CHI_TIEN_MAT_INPUT"]
    receipt_columns = {cell.value: cell.column for cell in receipt[1]}
    payment_columns = {cell.value: cell.column for cell in payment[1]}
    assert receipt.cell(2, receipt_columns["Mã ĐT"]).value is None
    assert payment.cell(2, payment_columns["Mã ĐT"]).value is None
    assert receipt.cell(2, receipt_columns["Người nộp tiền"]).value == "NGUYEN VAN AN"
    assert payment.cell(2, payment_columns["Người nhận tiền"]).value == "Lê Minh Tâm"
    assert receipt.cell(2, receipt_columns["Tỷ giá"]).value == 26120.125
    assert receipt.cell(2, receipt_columns["Thành tiền"]).value == 3265.015625

    exception = workbook["EXCEPTION"]
    exception_columns = {cell.value: cell.column for cell in exception[1]}
    assert exception.cell(2, exception_columns["Số tiền ngoại tệ"]).value == 3.75
    assert exception.cell(2, exception_columns["Tỷ giá"]).value is None
    assert exception.cell(2, exception_columns["Thành tiền"]).value is None

    summary = pd.read_excel(result.summary_path, sheet_name=SUMMARY_SHEET_NAME, dtype=object).fillna("")
    assert len(summary) == 5
    skipped = summary.loc[summary["Kết quả phân loại"] == "SKIPPED"].iloc[0]
    assert skipped[STATUS_COLUMN] == STATUS_DONE
    assert skipped[VOUCHER_COLUMN] == ""
    assert "BANK_FEE" in skipped[MESSAGE_COLUMN]
    assert "0.5 USD" in skipped[MESSAGE_COLUMN]
    assert result.stats["skipped_count"] == 1
    assert result.stats["exception_count"] == 1
    assert result.stats["auto_process_count"] == 3


def test_rtm_receipt_and_payment_can_be_completed_independently(tmp_path):
    processed = _processed_usd_batch()
    first = write_outputs(processed, tmp_path, {"output": {}})
    receipt_uid = next(item.transaction_uid for item in processed if item.flow == "thu_tien_mat")
    payment_uid = next(item.transaction_uid for item in processed if item.flow == "chi_tien_mat")
    assert receipt_uid != payment_uid

    mark_rpa_done(first.summary_path, receipt_uid, first.run_id, voucher_no="PT001")
    second = write_outputs(processed, tmp_path, {"output": {}})
    workbook = load_workbook(second.excel_path, data_only=True)
    assert workbook["THU_TIEN_MAT_INPUT"].max_row == 1
    assert workbook["CHI_TIEN_MAT_INPUT"].max_row == 2

    summary = pd.read_excel(second.summary_path, sheet_name=SUMMARY_SHEET_NAME, dtype=object).fillna("")
    rows = summary.set_index("transaction_uid")
    assert rows.at[receipt_uid, STATUS_COLUMN] == STATUS_DONE
    assert rows.at[receipt_uid, VOUCHER_COLUMN] == "PT001"
    assert rows.at[payment_uid, STATUS_COLUMN] != STATUS_DONE
